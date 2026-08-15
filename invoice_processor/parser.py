import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document

from .models import Invoice


FIELD_PATTERNS = {
    "invoice_number": r"(?:invoice\s*(?:number|no\.?|#))\s*[:\-]?\s*([A-Z0-9\-]+)",
    "vendor": r"(?:vendor|supplier)\s*[:\-]\s*(.+)",
    "invoice_date": r"(?:invoice\s*)?date\s*[:\-]\s*(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})",
    "subtotal": r"subtotal\s*[:\-]?\s*\$?([\d,]+(?:\.\d{1,2})?)",
    "tax": r"(?:tax|hst|gst)\s*[:\-]?\s*\$?([\d,]+(?:\.\d{1,2})?)",
    "total": r"(?<!sub)total\s*[:\-]?\s*\$?([\d,]+(?:\.\d{1,2})?)",
}


class InvoiceParseError(ValueError):
    pass


def read_document_text(path: Path) -> str:
    document = Document(path)
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(": ".join(values))
    return "\n".join(lines)


def parse_invoice_text(text: str, source_file: Path) -> Invoice:
    values = {}
    for name, pattern in FIELD_PATTERNS.items():
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        if not match:
            raise InvoiceParseError(f"Missing required field: {name.replace('_', ' ')}")
        values[name] = match.group(1).strip()

    try:
        invoice_date = parse_date(values["invoice_date"])
        subtotal = parse_money(values["subtotal"])
        tax = parse_money(values["tax"])
        total = parse_money(values["total"])
    except (ValueError, InvalidOperation) as exc:
        raise InvoiceParseError(str(exc)) from exc

    return Invoice(
        invoice_number=values["invoice_number"],
        vendor=values["vendor"],
        invoice_date=invoice_date,
        subtotal=subtotal,
        tax=tax,
        total=total,
        source_file=source_file,
    )


def parse_invoice(path: Path) -> Invoice:
    return parse_invoice_text(read_document_text(path), path)


def parse_money(value: str) -> Decimal:
    amount = Decimal(value.replace(",", "")).quantize(Decimal("0.01"))
    if amount < 0:
        raise ValueError("Amounts cannot be negative")
    return amount


def parse_date(value: str):
    normalized = value.replace("/", "-")
    formats = ("%Y-%m-%d", "%m-%d-%Y", "%d-%m-%Y")
    for date_format in formats:
        try:
            return datetime.strptime(normalized, date_format).date()
        except ValueError:
            continue
    raise ValueError(f"Unsupported invoice date: {value}")
