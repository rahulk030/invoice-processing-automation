from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Invoice, ProcessingIssue


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")


def write_excel_report(
    invoices: list[Invoice], issues: list[ProcessingIssue], output_path: Path
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    invoice_rows = [
        {
            "Invoice Number": invoice.invoice_number,
            "Vendor": invoice.vendor,
            "Invoice Date": invoice.invoice_date,
            "Subtotal": float(invoice.subtotal),
            "Tax": float(invoice.tax),
            "Total": float(invoice.total),
            "Calculated Total": float(invoice.calculated_total),
            "Difference": float(invoice.difference),
            "Source File": invoice.source_file.name,
        }
        for invoice in invoices
    ]
    issue_rows = [
        {
            "Source File": issue.source_file.name,
            "Severity": issue.severity.title(),
            "Message": issue.message,
        }
        for issue in issues
    ]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pd.DataFrame(invoice_rows, columns=[
            "Invoice Number", "Vendor", "Invoice Date", "Subtotal", "Tax", "Total",
            "Calculated Total", "Difference", "Source File"
        ]).to_excel(writer, sheet_name="Invoices", index=False)
        pd.DataFrame(issue_rows, columns=["Source File", "Severity", "Message"]).to_excel(
            writer, sheet_name="Exceptions", index=False
        )
        summary = pd.DataFrame(
            [
                {"Metric": "Invoices processed", "Value": len(invoices)},
                {"Metric": "Exceptions found", "Value": len(issues)},
                {"Metric": "Total invoice value", "Value": sum(float(i.total) for i in invoices)},
            ]
        )
        summary.to_excel(writer, sheet_name="Summary", index=False)

        for sheet in writer.book.worksheets:
            style_sheet(sheet)
        invoices_sheet = writer.book["Invoices"]
        for column in (4, 5, 6, 7, 8):
            for cell in invoices_sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                cell[0].number_format = '$#,##0.00'
        summary_sheet = writer.book["Summary"]
        summary_sheet["B4"].number_format = '$#,##0.00'
        for row in invoices_sheet.iter_rows(min_row=2):
            if row[7].value and abs(row[7].value) > 0.01:
                for cell in row:
                    cell.fill = WARNING_FILL


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for index, column_cells in enumerate(sheet.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 50)
        sheet.column_dimensions[get_column_letter(index)].width = width
