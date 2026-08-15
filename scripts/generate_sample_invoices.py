from decimal import Decimal
from pathlib import Path

from docx import Document


SAMPLES = [
    {
        "invoice_number": "INV-1042",
        "vendor": "North Shore Office Supply",
        "date": "2026-07-12",
        "subtotal": Decimal("245.00"),
        "tax": Decimal("31.85"),
        "total": Decimal("276.85"),
    },
    {
        "invoice_number": "INV-1043",
        "vendor": "Lakeside Packaging",
        "date": "2026-07-14",
        "subtotal": Decimal("180.00"),
        "tax": Decimal("23.40"),
        "total": Decimal("205.40"),  # Deliberate $2 exception for the report.
    },
]


def main() -> None:
    output_dir = Path(__file__).parents[1] / "examples" / "invoices"
    output_dir.mkdir(parents=True, exist_ok=True)
    for sample in SAMPLES:
        document = Document()
        document.add_heading("INVOICE", level=1)
        document.add_paragraph(f"Invoice Number: {sample['invoice_number']}")
        document.add_paragraph(f"Vendor: {sample['vendor']}")
        document.add_paragraph(f"Invoice Date: {sample['date']}")
        table = document.add_table(rows=3, cols=2)
        rows = [
            ("Subtotal", sample["subtotal"]),
            ("Tax", sample["tax"]),
            ("Total", sample["total"]),
        ]
        for row, (label, amount) in zip(table.rows, rows):
            row.cells[0].text = label
            row.cells[1].text = f"${amount:.2f}"
        document.save(output_dir / f"{sample['invoice_number']}.docx")
    print(f"Created {len(SAMPLES)} sample invoices in {output_dir}")


if __name__ == "__main__":
    main()
